"""Cria exemplo/dados_exemplo.xlsx: uma planilha com a mesma estrutura da
real, mas com proprietários, fornecedores e valores inventados.
É ela que roda na versão pública (portfólio) do painel."""
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

random.seed(42)
SAIDA = Path(__file__).resolve().parent / "dados_exemplo.xlsx"

PROPS = [
    ("Helena Duarte", "101", 0.10), ("Otávio Lins", "102", 0.10), ("Family Office Alfa", "201 e 202", 0.20),
    ("Beatriz Moura", "301", 0.10), ("Caio Ventura", "302", 0.10), ("Lúcia Prado", "401 e 402", 0.20),
    ("Rodrigo Sales", "501", 0.20),
]
FORN = [("Gerenciadora de Obra", 9000, 30000), ("Concessionária de Luz", 300, 1600), ("Água e Esgoto", 400, 1800),
        ("Assessoria Contábil", 600, 900), ("Escritório Jurídico", 2000, 12000), ("Segurança Patrimonial", 4000, 7000),
        ("Impostos", 800, 5000), ("Construtora Exemplo", 20000, 90000)]


def add_tabela(ws, nome, ref):
    t = Table(displayName=nome, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato"

    # Tabela1: proprietários
    ws.append([None, "Proprietarios", "Apartamentos", "Percentual"])
    ws.delete_rows(1)
    ws["B3"], ws["C3"], ws["D3"] = "Proprietarios", "Apartamentos", "Percentual"
    for i, (n, a, p) in enumerate(PROPS, start=4):
        ws.cell(i, 2, n), ws.cell(i, 3, a), ws.cell(i, 4, p)
    add_tabela(ws, "Tabela1", f"B3:D{3 + len(PROPS)}")

    ws["B16"], ws["C16"] = "Caixinha inicial G+P", 500

    # Tabela5: caixinha
    for j, h in enumerate(["Data", "Mês", "Fornecedor", "Descrição", "Despesa", "Saldo", "Receita"]):
        ws.cell(18, 2 + j, h)
    caix = [(dt.date(2026, 6, 12), "Aplicativo de transporte", "Visita à obra", 48.9),
            (dt.date(2026, 7, 3), "Cartório", "Autenticação de documentos", 112.0)]
    saldo = 500.0
    for k, (d, f, ds, v) in enumerate(caix, start=19):
        saldo -= v
        for j, val in enumerate([d, dt.date(d.year, d.month, 1), f, ds, v, saldo, 0]):
            ws.cell(k, 2 + j, val)
    add_tabela(ws, "Tabela5", f"B18:H{18 + len(caix)}")

    # RateioContrapartida
    for j, h in enumerate(["Apartamento", "Proprietário", "Contrapartida"]):
        ws.cell(25, 2 + j, h)
    for k, (n, a, p) in enumerate(PROPS, start=26):
        ws.cell(k, 2, a), ws.cell(k, 3, n), ws.cell(k, 4, round(random.uniform(60000, 140000), 2))
    add_tabela(ws, "RateioContrapartida", f"B25:D{25 + len(PROPS)}")

    # Tabela3: extrato do banco
    cab = ["Data", "Mês", "Proprietário", "Fornecedor", "receita", "despesa", "Saldo", "Descrição"]
    for j, h in enumerate(cab):
        ws.cell(3, 10 + j, h)
    linhas = [(dt.date(2025, 11, 10), None, None, None, None, "Saldo Anterior")]
    d = dt.date(2025, 11, 10)
    for mes in range(12):
        d = dt.date(2025 + (10 + mes) // 12, (10 + mes) % 12 + 1, 5)
        aporte = round(random.uniform(15000, 60000), -2)
        for n, a, p in PROPS:
            fator = 0.25 if n in ("Otávio Lins", "Caio Ventura") and mes >= 6 else 1.0  # dois proprietários atrasando
            linhas.append((d, n, None, round(aporte * p / 0.1 / 2 * fator, 2), None, f"Aporte {d:%m/%Y}"))
        for f, lo, hi in random.sample(FORN, 4):
            linhas.append((d.replace(day=random.randint(8, 27)), None, f, None, round(random.uniform(lo, hi), 2), f"{f} {d:%m/%Y}"))
        linhas.append((d.replace(day=28), None, "Rendimento Banco", round(random.uniform(0.1, 3), 2), None, "Rendimento conta"))
    contrap_total = sum(ws.cell(k, 4).value for k in range(26, 26 + len(PROPS)))
    linhas.append((dt.date(2026, 9, 10), None, "Contrapartida", None, round(contrap_total, 2), "Contrapartida"))
    for n, a, p in PROPS:
        linhas.append((dt.date(2026, 9, 15), n, None, round(contrap_total * p * (0.6 if n == "Otávio Lins" else 1.05), 2), None, "Aporte Contrapartida"))
    linhas.sort(key=lambda x: x[0])
    saldo = 0.0
    for k, (dd, prop, forn, rec, desp, desc) in enumerate(linhas, start=4):
        saldo += (rec or 0) - (desp or 0)
        for j, val in enumerate([dd, dt.date(dd.year, dd.month, 1), prop, forn, rec, desp, round(saldo, 2), desc]):
            ws.cell(k, 10 + j, val)
    add_tabela(ws, "Tabela3", f"J3:Q{3 + len(linhas)}")

    # debitosFuturos: projeções
    cab = ["Data", "Mês/ano", "Propretário", "Fornecedor", "receita", "despesa", "Saldo", "Descrição"]
    for j, h in enumerate(cab):
        ws.cell(3, 20 + j, h)
    fut = [(dt.date(2026, 11, 10), "Gerenciadora de Obra", 28000.0, "Gerenciamento novembro"),
           (dt.date(2026, 11, 20), "Impostos", 3900.0, "Retenções outubro"),
           (dt.date(2026, 11, 25), "Concessionária de Luz", 1450.0, "Luz novembro")]
    for k, (dd, f, v, ds) in enumerate(fut, start=4):
        for j, val in enumerate([dd, dt.date(dd.year, dd.month, 1), None, f, None, v, None, ds]):
            ws.cell(k, 20 + j, val)
    add_tabela(ws, "debitosFuturos", f"T3:AA{3 + len(fut)}")

    for c in "BJT":
        ws.column_dimensions[c].width = 14
    wb.save(SAIDA)
    print(f"Criado {SAIDA}")


if __name__ == "__main__":
    main()
